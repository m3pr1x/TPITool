"""Génère la packing list d'exemple livrée avec l'outil.

Données entièrement fictives. Le fichier reproduit volontairement les défauts
qu'on rencontre en vrai : en-tête sur deux lignes, dimensions éclatées sous un
titre fusionné, numéros de colis textuels, un colis étalé sur plusieurs lignes de
contenu, des cotes arrondies au demi-mètre, un volume déclaré faux, et deux colis
hors gabarit noyés au milieu des autres.

    python exemples/generer_exemple.py
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ICI = os.path.dirname(os.path.abspath(__file__))

COLIS = [
    # (n°, type, L, l, h, net, brut, volume déclaré, contenu)
    ("01  of  22", "Caisse bois", 120, 80, 96, 310, 380, 0.92, "Pompe centrifuge PC-101 — corps + roue"),
    ("2  of  22",  "Caisse bois", 120, 80, 96, 310, 380, 0.92, "Pompe centrifuge PC-102 — corps + roue"),
    ("3  of  22",  "Caisse bois", 244, 118, 104, 980, 1150, 2.99, "Échangeur E-201, faisceau tubulaire"),
    ("4  of  22",  "Caisse bois", 244, 118, 104, 980, 1150, 2.99, "Échangeur E-202, faisceau tubulaire"),
    ("5  of  22",  "Palette",     120, 100, 72, 420, 465, 0.86, "Robinetterie DN50 à DN150, lot complet"),
    ("6  of  22",  "Palette",     120, 100, 72, 400, 445, 0.86, "Robinetterie DN200, lot complet"),
    ("7  of  22",  "Caisse bois", 318, 96, 88, 1240, 1420, 2.69, "Tronçon de tuyauterie inox, ligne 12\""),
    ("8  of  22",  "Caisse bois", 318, 96, 88, 1240, 1420, 2.69, "Tronçon de tuyauterie inox, ligne 12\""),
    ("9  of  22",  "Caisse bois", 210, 140, 145, 1870, 2100, 4.26, "Réducteur RG-05 avec accouplement"),
    ("10  of  22", "Caisse bois", 165, 112, 118, 690, 810, 2.18, "Moteur électrique 75 kW"),
    ("11  of  22", "Caisse bois", 165, 112, 118, 690, 810, 2.18, "Moteur électrique 75 kW"),
    ("12  of  22", "Caisse bois", 100, 100, 100, 250, 300, 1.00, "Instrumentation — cotes à confirmer"),
    ("13  of  22", "Caisse bois", 150, 150, 150, 480, 560, 3.38, "Vannes de régulation — cotes à confirmer"),
    ("14  of  22", "Caisse bois", 286, 224, 196, 3100, 3480, 12.55, "Skid hydraulique complet"),
    ("15  of  22", "Caisse bois", 402, 168, 132, 2250, 2560, 8.91, "Charpente support, éléments longs"),
    ("16  of  22", "Caisse bois", 402, 168, 132, 2250, 2560, 3.20, "Charpente support, éléments longs"),
    ("17  of  22", "Caisse bois", 96, 64, 58, 145, 178, 0.36, "Pièces de rechange première année"),
    ("18  of  22", "Caisse bois", 96, 64, 58, 145, 178, 0.36, "Pièces de rechange deuxième année"),
    ("19  of  22", "Caisse bois", 268, 262, 312, 5400, 5980, 21.91, "Cuve tampon V-301, calorifugée"),
    ("20  of  22", "Caisse bois", 455, 244, 178, 4100, 4620, 19.76, "Sécheur d'air, module assemblé"),
    ("21  of  22", "Palette",     120, 80, 42, 180, 205, 0.40, "Documentation technique et outillage"),
    ("22  of  22", "Caisse bois", 178, 96, 84, 520, 615, 1.44, "Armoire électrique TGBT"),
]

# Colis 5 : plusieurs lignes de contenu, sans dimensions répétées.
SOUS_LIGNES = {
    "5  of  22": ["Robinets à tournant sphérique DN50, 12 pièces",
                  "Clapets anti-retour DN80, 6 pièces",
                  "Filtres en Y DN100, 4 pièces"],
}

GRIS = PatternFill("solid", fgColor="EFF3F4")
BLEU = PatternFill("solid", fgColor="D9E7EA")
BORD = Border(*[Side(style="thin", color="BFCCD0")] * 4)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing list"

    ws["A1"] = "PACKING LIST"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "N° PL-DEMO-2026-014"
    ws["A2"].font = Font(size=11, color="55666B")

    entetes = [
        ("A4", "EXPÉDITEUR"),      ("D4", "FOURNISSEUR"),      ("G4", "DESTINATAIRE"),
        ("A5", "Société Exemple"), ("D5", "Atelier Rhodanien"), ("G5", "Chantier Nord"),
        ("A6", "12 rue des Docks"),("D6", "ZI de la Plaine"),   ("G6", "Zone portuaire"),
        ("A7", "69007 Lyon"),      ("D7", "38070 Saint-Quentin"),("G7", "59140 Dunkerque"),
        ("A8", "FRANCE"),          ("D8", "FRANCE"),            ("G8", "FRANCE"),
    ]
    for cell, val in entetes:
        ws[cell] = val
        if cell[1] == "4":
            ws[cell].font = Font(bold=True, size=9, color="55666B")

    ws["A10"] = "MODE DE TRANSPORT"
    ws["A10"].font = Font(bold=True, size=9, color="55666B")
    ws["A11"] = "SEA — STACKABLE YES ON CONDITIONS (SAME BOX SIZES)"

    # En-tête du tableau : deux lignes, dimensions sous un titre fusionné.
    L = 13
    ws.cell(L, 1, "PACKAGE N°"); ws.cell(L, 2, "TYPE OF PACKAGE")
    ws.cell(L, 3, "MEASUREMENTS IN CM")
    ws.merge_cells(start_row=L, start_column=3, end_row=L, end_column=5)
    ws.cell(L, 6, "WEIGHT IN KG")
    ws.merge_cells(start_row=L, start_column=6, end_row=L, end_column=7)
    ws.cell(L, 8, "VOLUME"); ws.cell(L, 9, "QUANTITY"); ws.cell(L, 10, "PACKING CONTENT")

    ws.cell(L + 1, 3, "Length"); ws.cell(L + 1, 4, "Width"); ws.cell(L + 1, 5, "Height")
    ws.cell(L + 1, 6, "NET"); ws.cell(L + 1, 7, "GROSS"); ws.cell(L + 1, 8, "M3")

    for r in (L, L + 1):
        for c in range(1, 11):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True, size=9)
            cell.fill = BLEU if r == L else GRIS
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORD

    ligne = L + 2
    for (num, typ, lo, la, ha, net, brut, vol, contenu) in COLIS:
        ws.cell(ligne, 1, num); ws.cell(ligne, 2, typ)
        ws.cell(ligne, 3, lo);  ws.cell(ligne, 4, la); ws.cell(ligne, 5, ha)
        ws.cell(ligne, 6, net); ws.cell(ligne, 7, brut); ws.cell(ligne, 8, vol)
        ws.cell(ligne, 9, 1);   ws.cell(ligne, 10, contenu)
        for c in range(1, 11):
            ws.cell(ligne, c).border = BORD
        ligne += 1
        for extra in SOUS_LIGNES.get(num, []):
            ws.cell(ligne, 10, extra)
            for c in range(1, 11):
                ws.cell(ligne, c).border = BORD
            ligne += 1

    ligne += 1
    ws.cell(ligne, 2, "TOTAL").font = Font(bold=True)
    ws.cell(ligne, 6, sum(c[5] for c in COLIS)).font = Font(bold=True)
    ws.cell(ligne, 7, sum(c[6] for c in COLIS)).font = Font(bold=True)
    ws.cell(ligne, 8, round(sum(c[2] * c[3] * c[4] / 1e6 for c in COLIS), 2)).font = Font(bold=True)
    ws.cell(ligne, 9, len(COLIS)).font = Font(bold=True)

    for col, w in zip("ABCDEFGHIJ", (14, 16, 10, 10, 10, 10, 10, 10, 10, 46)):
        ws.column_dimensions[col].width = w

    sortie = os.path.join(ICI, "packing-list-exemple.xlsx")
    wb.save(sortie)
    vol = sum(c[2] * c[3] * c[4] / 1e6 for c in COLIS)
    print(f"{sortie}\n{len(COLIS)} colis · {vol:.1f} m³ · {sum(c[6] for c in COLIS)} kg")


if __name__ == "__main__":
    main()
